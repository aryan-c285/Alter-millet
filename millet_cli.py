#!/usr/bin/env python3

import os, sys, glob, re, difflib, argparse
import pandas as pd
import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
NUTKEYS = ['kcal', 'carb', 'protein', 'fat', 'fibre']
NUTLABEL = {'kcal': 'Energy(kcal)', 'carb': 'Carb(g)', 'protein': 'Protein(g)',
            'fat': 'Fat(g)', 'fibre': 'Fibre(g)'}

# -----------------------------------------------------------------------------
# file discovery (tolerant of spaces vs underscores in names)
# -----------------------------------------------------------------------------
def find(patterns, label):
    for pat in patterns:
        hits = glob.glob(os.path.join(HERE, pat))
        if hits:
            return hits[0]
    sys.exit("ERROR: could not find the %s in this folder:\n  %s\n  looked for: %s\n  present: %s"
             % (label, HERE, ', '.join(patterns),
                [f for f in os.listdir(HERE) if f.lower().endswith(('.csv', '.xlsx'))]))

MAP_FILE = find(['*Mapping*Final*.xlsx', '*Millet_Mapping*.xlsx'], 'colour-coded mapping (Millet_Mapping_Final_1.xlsx)')
NUT_FILE = find(['*Nutrition*100*.csv', '*Nutrition*.csv', '*utrition*.csv'], 'nutrition CSV')
MIL_FILE = find(['*155*.xlsx', '*Final*.xlsx'], 'millet recipe file (155_Final.xlsx)')

# -----------------------------------------------------------------------------
# 1. NUTRITION TABLE  (per 100 g)
# -----------------------------------------------------------------------------
def _num(v):
    """Parse a possibly-ranged / non-numeric cell to a float (range -> midpoint)."""
    s = str(v).strip().replace('–', '-')
    if s in ('', '-', 'nan') or 'source' in s.lower():
        return np.nan
    m = re.findall(r'\d+\.?\d*', s)
    return round(sum(map(float, m)) / len(m), 2) if m else np.nan

def load_nutrition():
    nut = pd.read_csv(NUT_FILE)
    colmap = {'Energy (kcal/100g)': 'kcal', 'Carbohydrates (g/100g)': 'carb',
              'Protein (g/100g)': 'protein', 'Fat (g/100g)': 'fat',
              'Dietary Fibre (g/100g)': 'fibre'}
    for c, k in colmap.items():
        nut[k] = nut[c].map(_num)
    by_id = nut.set_index('ID')
    # name -> per-100g, for fuzzy ingredient lookup
    name_table = []
    for _, r in nut.iterrows():
        for nm in {str(r['Standard Ingredient']).lower(), str(r['Form / Variety']).lower()}:
            if nm and nm != 'nan':
                vals = {k: (None if pd.isna(r[k]) else float(r[k])) for k in NUTKEYS}
                name_table.append((nm, vals, str(r['Standard Ingredient'])))
    return nut, by_id, name_table

NUT, NUT_BY_ID, NAME_TABLE = load_nutrition()

def grain_nutrition(name):
    """Per-100g nutrition for a replaced ingredient name, or None if not matched."""
    if not isinstance(name, str) or not name.strip():
        return None
    nm = name.lower()
    best, score, label = None, 0, None
    for cand, vals, lab in NAME_TABLE:
        sc = difflib.SequenceMatcher(None, nm, cand).ratio()
        if nm in cand or cand in nm:
            sc = max(sc, 0.9)
        if sc > score:
            score, best, label = sc, vals, lab
    return (dict(best, matched_as=label)) if score >= 0.6 else None

# -----------------------------------------------------------------------------
# 2. MILLET NUTRITION  (per 100 g; Multi Millet derived from its blend)
# -----------------------------------------------------------------------------
BASE_MILLET = {'Sorghum': 'A005', 'Pearl Millet': 'A003', 'Finger Millet': 'A010',
               'Little Millet': 'A016', 'Kodo Millet': 'A017', 'Proso Millet': 'NEW-01',
               'Foxtail Millet': 'NEW-03', 'Barnyard Millet': 'NEW-04'}
NAME2TYPE = {'sorghum': 'Sorghum', 'finger millet': 'Finger Millet', 'pearl millet': 'Pearl Millet',
             'foxtail': 'Foxtail Millet', 'little millet': 'Little Millet', 'kodo millet': 'Kodo Millet',
             'proso millet': 'Proso Millet', 'barnyard millet': 'Barnyard Millet'}

def _base(t):
    r = NUT_BY_ID.loc[BASE_MILLET[t]]
    return {k: round(float(r[k]), 2) for k in NUTKEYS}

def load_millet_recipes():
    mil = pd.read_excel(MIL_FILE)
    return dict(zip(mil.recipe_name, mil.ingredients_raw)), dict(zip(mil.recipe_name, mil.millet_type))
RAW_OF, TYPE_OF = load_millet_recipes()

def millet_nutrition(recipe_name, millet_type):
    """Per-100g nutrition for a millet recipe; blends are averaged from components."""
    t = millet_type or TYPE_OF.get(recipe_name)
    if t and t != 'Multi Millet':
        return dict(_base(t), src=t)
    raw = RAW_OF.get(recipe_name, '')
    m = re.search(r'\(([^)]*)\)', str(raw))
    comps = []
    if m:
        for nm, ty in NAME2TYPE.items():
            if nm in m.group(1).lower() and ty not in comps:
                comps.append(ty)
    if not comps:
        comps = ['Sorghum', 'Finger Millet', 'Pearl Millet', 'Foxtail Millet']
    return dict({k: round(sum(_base(c)[k] for c in comps) / len(comps), 2) for k in NUTKEYS},
                src=' + '.join(comps))

# -----------------------------------------------------------------------------
# 3. COLOUR-CODED MAPPING  (read cell fills straight from the Excel)
# -----------------------------------------------------------------------------
COLOUR_STATUS = {
    'FF00FF00': 'ALTERNATIVE', 'FFFFE599': 'ALTERNATIVE', '00000000': 'ALTERNATIVE', None: 'ALTERNATIVE',
    'FF4A86E8': 'RECOMMENDATION', 'FF00FFFF': 'RECOMMENDATION',
    'FFFF0000': 'EXCLUDE', 'FFE06666': 'EXCLUDE', 'FFEA9999': 'EXCLUDE',
}

def load_mapping():
    """Return the kept pairings (ALTERNATIVE/RECOMMENDATION) from the coloured sheet."""
    wb = openpyxl.load_workbook(MAP_FILE)
    ws = wb['Top 3 Candidates']
    hdr = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr)}
    pairs = []
    for ri in range(2, ws.max_row + 1):
        rgb = getattr(ws.cell(ri, 1).fill.fgColor, 'rgb', None)   # colour on millet-id cell
        status = COLOUR_STATUS.get(rgb, 'ALTERNATIVE')
        if status == 'EXCLUDE':
            continue
        row = {h: ws.cell(ri, idx[h] + 1).value for h in hdr}
        oid = row.get('original_recipe_id')
        pairs.append(dict(
            original_id=str(int(oid)) if isinstance(oid, (int, float)) and not pd.isna(oid) else str(oid),
            original=row.get('original_recipe'),
            food_family=row.get('food_family'),
            millet_recipe=row.get('millet_recipe'),
            millet_type=row.get('millet_type'),
            replaced=row.get('replaced_ingredient'),
            status=status))
    return pairs

MAPPING = load_mapping()

# -----------------------------------------------------------------------------
# 4. DISEASE RULES  (+1 = more of this nutrient is better; -1 = less is better)
#    encoded from the condition nutrition matrices
# -----------------------------------------------------------------------------
DISEASES = {
    'Type 2 Diabetes':        {'fibre': 1, 'carb': -1, 'fat': -1},
    'High Cholesterol':       {'fibre': 1, 'fat': -1, 'kcal': -1},
    'Cardiovascular Disease': {'fibre': 1, 'fat': -1},
    'Hypertension':           {'fibre': 1, 'kcal': -1},
    'Chronic Kidney Disease': {'protein': -1},
    'MASLD / Fatty Liver':    {'kcal': -1, 'carb': -1, 'fat': -1, 'fibre': 1},
    'Obesity / Overweight':   {'kcal': -1, 'fibre': 1},
    'Gout':                   {'carb': -1, 'fat': -1, 'fibre': 1, 'protein': -1},
    'Osteoporosis':           {'protein': 1},
    'Dental Caries':          {},   # driven by sugar + calcium: not in ingredient data
}

def resolve_disease(text):
    t = text.lower()
    for name in DISEASES:
        if t in name.lower() or name.lower().startswith(t):
            return name
    hits = difflib.get_close_matches(text, list(DISEASES), n=1, cutoff=0.4)
    return hits[0] if hits else None

# -----------------------------------------------------------------------------
# 5. COMPARISON  (millet vs the marked replaced grain, per 100 g)
# -----------------------------------------------------------------------------
def compare(pair, disease):
    rule = DISEASES[disease]
    mn = millet_nutrition(pair['millet_recipe'], pair['millet_type'])
    gn = grain_nutrition(pair['replaced']) if pair['replaced'] else None
    if gn is None:
        return dict(verdict='RECOMMENDATION (no grain marked to compare)', lines=[], millet_src=mn['src'], grain=None)
    if not rule:
        return dict(verdict='NOT COMPUTABLE (condition driven by nutrients not in data)',
                    lines=[], millet_src=mn['src'], grain=gn.get('matched_as'))
    lines, score = [], 0
    for k in NUTKEYS:
        in_rule = k in rule
        dv = round(mn[k] - gn[k], 1)
        if in_rule:
            good = dv * rule[k] > 0
            score += 1 if good else (-1 if dv * rule[k] < 0 else 0)
            mark = '+' if good else ('x' if dv * rule[k] < 0 else '=')
            want = 'want higher' if rule[k] > 0 else 'want lower'
        else:
            mark, want = '.', 'not weighted'
        lines.append((NUTLABEL[k], mn[k], gn[k], dv, want, mark))
    verdict = 'RECOMMEND millet (better for this condition)' if score > 0 else \
              ('KEEP ORIGINAL (millet worse here)' if score < 0 else 'NO CLEAR GAIN')
    return dict(verdict=verdict, lines=lines, millet_src=mn['src'], grain=gn.get('matched_as'))

# -----------------------------------------------------------------------------
# 6. CLI
# -----------------------------------------------------------------------------
def list_foods():
    names = sorted({p['original'] for p in MAPPING if p['original']})
    for n in names:
        print(' ', n)
    print('\n%d dishes available.' % len(names))

def list_diseases():
    for d in DISEASES:
        note = '  (not computable from current nutrients)' if not DISEASES[d] else ''
        print(' ', d, note)

def run(food, disease):
    dname = resolve_disease(disease)
    if not dname:
        print("Unknown condition '%s'. Try: python millet_cli.py --list-diseases" % disease)
        return
    mine = [p for p in MAPPING if p['original'] and food.lower() in p['original'].lower()]
    if not mine:
        print("No dish matches '%s'. Try: python millet_cli.py --list-foods" % food)
        return
    exact = sorted({p['original'] for p in mine})
    target = exact[0]
    mine = [p for p in mine if p['original'] == target]
    alts = [p for p in mine if p['status'] == 'ALTERNATIVE']
    recs = [p for p in mine if p['status'] == 'RECOMMENDATION']

    print('\n' + '=' * 70)
    print('DISH : %s' % target)
    print('FOR  : %s' % dname)
    print('=' * 70)
    for label, group in [('ALTERNATIVES (mapped swaps)', alts),
                         ('RECOMMENDATIONS (weaker / standalone)', recs)]:
        if not group:
            continue
        print('\n--- %s ---' % label)
        for p in group:
            c = compare(p, dname)
            print('\n  * %s  [%s]' % (p['millet_recipe'], p['millet_type']))
            if p['replaced']:
                print('    millet (%s) replaces: %s' % (c['millet_src'], p['replaced']))
            print('    VERDICT: %s' % c['verdict'])
            if c['lines']:
                print('    %-13s %8s %8s %8s   %s' % ('nutrient/100g', 'millet', 'grain', 'delta', 'direction'))
                for nm, mv, gv, dv, want, mark in c['lines']:
                    print('    %-13s %8.1f %8.1f %+8.1f   %s %s' % (nm, mv, gv, dv, mark, want))
    print()

def interactive():
    print('\nMILLET ALTERNATIVE FINDER  (type a condition, then a dish)\n')
    list_diseases()
    while True:
        d = input('\nCondition (or "q" to quit): ').strip()
        if d.lower() in ('q', 'quit', 'exit'):
            break
        dname = resolve_disease(d)
        if not dname:
            print("  not recognised — pick from the list above")
            continue
        food = input('Dish you\'d normally eat (partial name ok): ').strip()
        if food:
            run(food, dname)

def main():
    ap = argparse.ArgumentParser(description='Millet alternative finder (CLI)')
    ap.add_argument('--food', help='original dish name (partial ok)')
    ap.add_argument('--disease', help='health condition (partial ok)')
    ap.add_argument('--list-foods', action='store_true')
    ap.add_argument('--list-diseases', action='store_true')
    a = ap.parse_args()
    print('Using mapping   : %s' % os.path.basename(MAP_FILE))
    print('Using nutrition : %s' % os.path.basename(NUT_FILE))
    print('Using millets   : %s' % os.path.basename(MIL_FILE))
    print('Loaded %d usable pairings (excluded rows already dropped).' % len(MAPPING))
    if a.list_foods:
        list_foods()
    elif a.list_diseases:
        list_diseases()
    elif a.food and a.disease:
        run(a.food, a.disease)
    else:
        interactive()

if __name__ == '__main__':
    main()
